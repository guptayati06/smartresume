import io
import re
import os

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import pytesseract
    from pdf2image import convert_from_bytes
    HAS_OCR = True
except Exception:
    HAS_OCR = False


def extract_text_from_bytes(filename: str, data: bytes) -> tuple[str, str, list[str]]:
    """
    Extracts text from PDF, DOCX, or TXT file bytes.
    Returns:
        (extracted_text, extraction_method, warnings)
    """
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    warnings = []
    method = "unknown"
    text = ""

    if suffix == ".pdf":
        method = "PDF text layer"
        # Try PyMuPDF first
        if fitz is not None:
            try:
                with fitz.open(stream=data, filetype="pdf") as pdf:
                    parts = []
                    for page in pdf:
                        page_text = page.get_text("text") or ""
                        parts.append(page_text)
                    text = "\n".join(parts)
            except Exception as e:
                warnings.append(f"PyMuPDF failed: {e}")

        # Fallback to pypdf if text is empty and pypdf available
        if (not text or len(text.strip()) < 10) and PdfReader is not None:
            try:
                reader = PdfReader(io.BytesIO(data))
                parts = [p.extract_text() or "" for p in reader.pages]
                text = "\n".join(parts)
                method = "PyPDF text layer"
            except Exception as e:
                warnings.append(f"PyPDF fallback error: {e}")

        # Check for scanned PDF / sparse text layer (< 30 non-whitespace characters)
        clean_len = len(re.sub(r"\s+", "", text))
        if clean_len < 30:
            if HAS_OCR:
                try:
                    pages = convert_from_bytes(data, dpi=250)
                    ocr_text = "\n".join(pytesseract.image_to_string(page, config="--psm 3") for page in pages)
                    if ocr_text.strip():
                        text = ocr_text
                        method = "OCR fallback"
                    else:
                        warnings.append("OCR completed but detected minimal text.")
                except Exception as ocr_err:
                    warnings.append(f"OCR fallback error: {ocr_err}")
            else:
                method = "Sparse PDF layer (OCR unavailable)"
                warnings.append("Text layer is sparse (<30 characters). Scanned document detected. Server OCR engine not active.")

    elif suffix == ".docx":
        method = "DOCX parser"
        if Document is not None:
            try:
                doc = Document(io.BytesIO(data))
                parts = [p.text for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        parts.append(" | ".join(cell.text.strip() for cell in row.cells))
                text = "\n".join(parts)
            except Exception as e:
                warnings.append(f"DOCX extraction error: {e}")
        else:
            warnings.append("python-docx not installed.")

    elif suffix == ".txt":
        method = "TXT parser"
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = data.decode("latin-1")
                warnings.append("Decoded using latin-1 encoding fallback.")
            except Exception as e:
                text = data.decode("utf-8", errors="replace")
                warnings.append("Decoded with replacement characters.")

    else:
        raise ValueError(f"Unsupported resume type for '{filename}'. Please upload a PDF, DOCX, or TXT document.")

    if not text.strip():
        raise ValueError(f"No readable text extracted from '{filename}'. Please check file formatting or provide a searchable document.")

    return text, method, warnings
